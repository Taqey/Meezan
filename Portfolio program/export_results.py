#!/usr/bin/env python3
"""Build the 'Results' workbook from computed portfolio results.

The browser app cannot run openpyxl, so the flow is two steps:
  1. In the app: result screen -> "JSON" button -> downloads results_<date>.json
     (produced by buildResultsPayload() in index.html — tickers, weights,
     per-stock expected returns and metrics of the current run, nothing hardcoded).
  2. Here: python export_results.py results_<date>.json [output.xlsx]

Output: one sheet named "Results" with two side-by-side tables
  LEFT  (A:D): Constant Correlation Model, English only, LTR.
  RIGHT (F:I): Single Index Model, Arabic only (tickers stay as-is).
  Column E is an empty spacer. Everything (n, k per model, tickers, weights,
metrics) comes from the payload.

Requires: pip install openpyxl
"""

import json
import sys
from copy import copy

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

GRAY_FILL = PatternFill("solid", fgColor="D9D9D9")
THIN = Side(style="thin", color="000000")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

BASE_FONT = Font(name="Calibri", size=11)
BOLD_FONT = Font(name="Calibri", size=11, bold=True)
TITLE_FONT = Font(name="Calibri", size=14, bold=True)
ITALIC_FONT = Font(name="Calibri", size=11, italic=True)

FMT_PCT = "0.00%"
FMT_3DP = "0.000"
FMT_2DP = "0.00"
FMT_AMT = "#,##0.00"
FMT_BLANK_ZERO = '#,##0.00;-#,##0.00;""'

COLUMN_WIDTHS = {"A": 34, "B": 14, "C": 16, "D": 18, "E": 4,
                 "F": 34, "G": 14, "H": 16, "I": 18}


def _style(cell, font=BASE_FONT, alignment=None, number_format=None,
           fill=None, border=None):
    cell.font = copy(font)
    if alignment is not None:
        cell.alignment = copy(alignment)
    if number_format is not None:
        cell.number_format = number_format
    if fill is not None:
        cell.fill = copy(fill)
    if border is not None:
        cell.border = copy(border)
    return cell


def _merge_bordered(ws, row, c0, c1, value, font=BASE_FONT, alignment=None,
                    number_format=None, fill=None):
    """Write value into merged row/c0:c1 with border on every cell."""
    for c in range(c0, c1 + 1):
        cell = ws.cell(row=row, column=c)
        cell.border = copy(BORDER)
        cell.font = copy(font)
        if alignment is not None:
            cell.alignment = copy(alignment)
        if number_format is not None:
            cell.number_format = number_format
    top = ws.cell(row=row, column=c0)
    top.value = value
    if fill is not None:
        for c in range(c0, c1 + 1):
            ws.cell(row=row, column=c).fill = copy(fill)
    ws.merge_cells(start_row=row, start_column=c0, end_row=row, end_column=c1)
    return top


def _bordered_block(ws, c0, c3, rows):
    for r in rows:
        for c in range(c0, c3 + 1):
            ws.cell(row=r, column=c).border = copy(BORDER)


def _fin(x, default=None):
    try:
        v = float(x)
    except (TypeError, ValueError):
        return default
    return v if v == v else default  # reject NaN

def build_results_workbook(payload):
    """payload: dict with keys n, rfAnn (percent), ccm|None, sim|None.

    Model dict: {stocks:[{ticker, weight, exAnn}]} where exAnn is the
    annualized expected EXCESS return as a fraction (None if unknown).
    CCM adds rho, sharpeAnn; SIM adds beta, treynorAnn (informational only —
    sheet formulas recompute treynor/total from the cells).
    """
    n = int(payload["n"])
    rf = float(payload["rfAnn"]) / 100.0
    ccm = payload.get("ccm")
    sim = payload.get("sim")

    ccm_stocks = sorted((ccm or {}).get("stocks", []),
                        key=lambda s: -s["weight"]) if ccm else []
    sim_stocks = sorted((sim or {}).get("stocks", []),
                        key=lambda s: -s["weight"]) if sim else []
    kc, ks = len(ccm_stocks), len(sim_stocks)
    kmax = max(kc, ks)

    wb = Workbook()
    ws = wb.active
    ws.title = "Results"
    ws.sheet_view.showGridLines = False
    ws.sheet_view.rightToLeft = False
    ws.sheet_view.zoomScale = 100

    for col, w in COLUMN_WIDTHS.items():
        ws.column_dimensions[col].width = w

    tables = [
        {"cols": (1, 4), "stocks": ccm_stocks, "k": kc,
         "title": "Constant Correlation Model",
         "subtitle": "Invest in {k} out of {n} stocks",
         "headers": ["Stock", "Weight (%)", "Amount", "Total Return"],
         "total": "Total", "rf_label": "Risk-free rate (T-bill)",
         "text_align": "left", "model": ccm},
        {"cols": (6, 9), "stocks": sim_stocks, "k": ks,
         "title": "نموذج المؤشر الواحد",
         "subtitle": "استثمر في {k} من أصل {n} أسهم",
         "headers": ["السهم", "الوزن (%)", "المبلغ", "العائد الكلي"],
         "total": "الإجمالي",
         "rf_label": "معدل العائد الخالي من المخاطر (أذون الخزانة)",
         "text_align": "right", "model": sim},
    ]

    LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
    RIGHT = Alignment(horizontal="right", vertical="center", wrap_text=True)
    CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
    NUM = Alignment(horizontal="right", vertical="center")

    for t in tables:
        c0, c3 = t["cols"]
        cT, cW, cA, cR = c0, c0 + 1, c0 + 2, c0 + 3
        k = t["k"]
        t_align = RIGHT if t["text_align"] == "right" else LEFT

        title = ws.cell(row=1, column=c0, value=t["title"])
        _style(title, font=TITLE_FONT, alignment=t_align)
        ws.merge_cells(start_row=1, start_column=c0, end_row=1, end_column=c3)
        sub = ws.cell(row=2, column=c0,
                      value=t["subtitle"].format(k=k, n=n))
        _style(sub, font=ITALIC_FONT, alignment=t_align)
        ws.merge_cells(start_row=2, start_column=c0, end_row=2, end_column=c3)
        ws.row_dimensions[3].height = 6

        for c, h in zip((cT, cW, cA, cR), t["headers"]):
            _style(ws.cell(row=4, column=c, value=h),
                   font=BOLD_FONT, alignment=CENTER,
                   fill=GRAY_FILL, border=BORDER)

        a_col = {1: "C", 6: "H"}[c0]
        for i, s in enumerate(t["stocks"]):
            r = 5 + i
            ex = _fin(s.get("exAnn"))
            _style(ws.cell(row=r, column=cT, value=s["ticker"]),
                   alignment=t_align)
            _style(ws.cell(row=r, column=cW, value=float(s["weight"])),
                   alignment=NUM, number_format=FMT_PCT, border=BORDER)
            _style(ws.cell(row=r, column=cA), alignment=NUM,
                   number_format=FMT_AMT, border=BORDER)
            if ex is None:
                _style(ws.cell(row=r, column=cR), alignment=NUM,
                       border=BORDER)
            else:
                _style(ws.cell(
                    row=r, column=cR,
                    value=f"=IF({a_col}{r}=\"\","
                          f"\"\",{a_col}{r}*(1+{ex + rf!r}))"),
                       alignment=NUM, number_format=FMT_AMT, border=BORDER)
        _bordered_block(ws, c0, c3, range(5, 5 + k))

        trow = 5 + k
        _style(ws.cell(row=trow, column=cT, value=t["total"]),
               font=BOLD_FONT, alignment=t_align,
               fill=GRAY_FILL, border=BORDER)
        if k > 0:
            wcol = {1: "B", 6: "G"}[c0]
            _style(ws.cell(row=trow, column=cW,
                            value=f"=SUM({wcol}5:{wcol}{4 + k})"),
                   font=BOLD_FONT, alignment=NUM, number_format=FMT_PCT,
                   fill=GRAY_FILL, border=BORDER)
            _style(ws.cell(row=trow, column=cA,
                            value=f"=SUM({a_col}5:{a_col}{4 + k})"),
                   font=BOLD_FONT, alignment=NUM, number_format=FMT_BLANK_ZERO,
                   fill=GRAY_FILL, border=BORDER)
            rcol = {1: "D", 6: "I"}[c0]
            _style(ws.cell(row=trow, column=cR,
                            value=f"=SUM({rcol}5:{rcol}{4 + k})"),
                   font=BOLD_FONT, alignment=NUM, number_format=FMT_BLANK_ZERO,
                   fill=GRAY_FILL, border=BORDER)
        else:
            for c in (cW, cA, cR):
                _style(ws.cell(row=trow, column=c),
                       font=BOLD_FONT, alignment=NUM,
                       fill=GRAY_FILL, border=BORDER)

        rfrow = 6 + k
        _style(ws.cell(row=rfrow, column=cT, value=t["rf_label"]),
               alignment=t_align)
        _style(ws.cell(row=rfrow, column=cW, value=rf),
               alignment=NUM, number_format=FMT_PCT)

    mhead = 9 + kmax
    rfC, rfS = 6 + kc, 6 + ks
    cEx = _fin((ccm or {}).get("exAnn"))
    sEx = _fin((sim or {}).get("exAnn"))

    _style(ws.cell(row=mhead, column=1, value="Metric"),
           font=BOLD_FONT, alignment=CENTER, fill=GRAY_FILL, border=BORDER)
    _style(ws.cell(row=mhead, column=2, value="Value"),
           font=BOLD_FONT, alignment=CENTER, fill=GRAY_FILL, border=BORDER)
    _merge_bordered(ws, mhead, 3, 4, "Description",
                    font=BOLD_FONT, alignment=CENTER, fill=GRAY_FILL)
    _style(ws.cell(row=mhead, column=6, value="المقياس"),
           font=BOLD_FONT, alignment=CENTER, fill=GRAY_FILL, border=BORDER)
    _style(ws.cell(row=mhead, column=7, value="القيمة"),
           font=BOLD_FONT, alignment=CENTER, fill=GRAY_FILL, border=BORDER)
    _merge_bordered(ws, mhead, 8, 9, "الوصف",
                    font=BOLD_FONT, alignment=CENTER, fill=GRAY_FILL)

    def metric(row, c0, name, value, fmt, desc, name_align):
        _style(ws.cell(row=row, column=c0, value=name),
               alignment=name_align, border=BORDER)
        vcell = ws.cell(row=row, column=c0 + 1)
        if isinstance(value, str) and value.startswith("="):
            vcell.value = value
        else:
            vcell.value = value if value is not None else "n/a"
        _style(vcell, alignment=NUM, number_format=fmt, border=BORDER)
        ws.row_dimensions[row].height = 32
        return vcell.coordinate

    r1, r2, r3, r4 = mhead + 1, mhead + 2, mhead + 3, mhead + 4
    rho = _fin((ccm or {}).get("rho"))
    beta = _fin((sim or {}).get("beta"))

    metric(r1, 1, "Average correlation", rho, FMT_3DP,
           "Same correlation assumed for every pair", LEFT)
    _merge_bordered(ws, r1, 3, 4,
                    "Same correlation assumed for every pair",
                    alignment=LEFT)
    metric(r1, 6, "بيتا المحفظة", beta, FMT_2DP,
           f"تتحرك بنسبة {beta:.2f}% مقابل كل 1% يتحركه السوق"
           if beta is not None else "n/a", RIGHT)
    _merge_bordered(
        ws, r1, 8, 9,
        f"تتحرك بنسبة {beta:.2f}% مقابل كل 1% يتحركه السوق"
        if beta is not None else "n/a", alignment=RIGHT)

    metric(r2, 1, "Expected excess return per year", cEx, FMT_PCT,
           "Extra return above the risk-free rate (T-bill), per year", LEFT)
    _merge_bordered(ws, r2, 3, 4,
                    "Extra return above the risk-free rate (T-bill), per year",
                    alignment=LEFT)
    metric(r2, 6, "العائد الزائد المتوقع سنوياً", sEx, FMT_PCT,
           "العائد الزائد فوق معدل أذون الخزانة، سنوياً", RIGHT)
    _merge_bordered(ws, r2, 8, 9,
                    "العائد الزائد فوق معدل أذون الخزانة، سنوياً",
                    alignment=RIGHT)

    cSh = _fin((ccm or {}).get("sharpeAnn"))
    metric(r3, 1, "Return per unit of risk", cSh, FMT_3DP,
           "Extra return for each unit of portfolio risk", LEFT)
    _merge_bordered(ws, r3, 3, 4,
                    "Extra return for each unit of portfolio risk",
                    alignment=LEFT)
    treynor = (f"=G{r2}/G{r1}"
               if sEx is not None and beta not in (None, 0) else "n/a")
    metric(r3, 6, "نسبة ترينور (سنوية)", treynor, FMT_3DP,
           "العائد الإضافي مقابل كل وحدة من مخاطر السوق", RIGHT)
    _merge_bordered(ws, r3, 8, 9,
                    "العائد الإضافي مقابل كل وحدة من مخاطر السوق",
                    alignment=RIGHT)

    tot_c = f"=B{r2}+$B${rfC}" if cEx is not None else "n/a"
    metric(r4, 1, "Total expected return", tot_c, FMT_PCT,
           "Total annual return including the T-bill rate", LEFT)
    _merge_bordered(ws, r4, 3, 4,
                    "Total annual return including the T-bill rate",
                    alignment=LEFT)
    tot_s = f"=G{r2}+$G${rfS}" if sEx is not None else "n/a"
    metric(r4, 6, "العائد الكلي المتوقع", tot_s, FMT_PCT,
           "العائد السنوي الكلي شامل عائد أذون الخزانة", RIGHT)
    _merge_bordered(ws, r4, 8, 9,
                    "العائد السنوي الكلي شامل عائد أذون الخزانة",
                    alignment=RIGHT)

    # Fresh sheets default to active cell A1, zoom 100%, nothing frozen.
    return wb


def main(argv):
    if len(argv) < 2 or len(argv) > 3:
        print(__doc__)
        return 2
    with open(argv[1], encoding="utf-8") as fh:
        payload = json.load(fh)
    wb = build_results_workbook(payload)
    out = argv[2] if len(argv) == 3 else "Results.xlsx"
    wb.save(out)
    print(f"wrote {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))


